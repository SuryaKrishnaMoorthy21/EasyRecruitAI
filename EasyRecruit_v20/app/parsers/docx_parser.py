"""
DOCX Parser - Extract text from Word documents using python-docx
"""
from docx import Document
from typing import Optional
import logging
import os
import time
import zipfile

from app.parsers.errors import (
    DocumentError, DocumentTimeoutError, DocumentTooComplexError, EncryptedDocumentError,
    MalformedDocumentError, NoExtractableTextError,
)

logger = logging.getLogger(__name__)

_DOCX_DAMAGED = "This file is not a valid DOCX document or it is damaged."
_XML_ENTRY_SUFFIXES = (".xml", ".rels")


def inspect_docx_package(
    path: str,
    *,
    max_uncompressed_bytes: int,
    max_entries: int = 1000,
    max_xml_part_bytes: Optional[int] = None,
) -> None:
    """
    Cheap structural check of an UNTRUSTED .docx, run BEFORE python-docx loads it (Task 5).

    A .docx is a ZIP and the upload limit only bounds its COMPRESSED size: a few hundred KB can expand
    to gigabytes of XML (a decompression bomb). The sizes declared in the central directory are what
    zipfile enforces while reading, so summing them bounds what python-docx can be made to load.
    Also rejects encrypted entries and ZIP archives that are not Word documents.
    """
    if max_xml_part_bytes is None:
        max_xml_part_bytes = max(1024 * 1024, max_uncompressed_bytes // 4)
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise MalformedDocumentError(_DOCX_DAMAGED, detail_for_log="BadZipFile") from None
    except Exception as exc:
        raise MalformedDocumentError(_DOCX_DAMAGED, detail_for_log=f"{type(exc).__name__}: {exc}") from None

    limit_mb = f"{max_uncompressed_bytes / (1024 * 1024):g} MB"
    with archive:
        infos = archive.infolist()
        if len(infos) > max_entries:
            raise DocumentTooComplexError(
                "This DOCX contains an unusual number of internal parts and cannot be processed."
            )
        total, names = 0, set()
        for info in infos:
            if info.flag_bits & 0x1:
                raise EncryptedDocumentError(
                    "This DOCX is password-protected. Please remove the password and upload it again."
                )
            total += info.file_size
            if total > max_uncompressed_bytes:
                raise DocumentTooComplexError(
                    f"This DOCX expands to more than {limit_mb} when opened, which is far more than a "
                    "resume needs, so it was not processed."
                )
            if info.filename.lower().endswith(_XML_ENTRY_SUFFIXES) and info.file_size > max_xml_part_bytes:
                raise DocumentTooComplexError(
                    "This DOCX contains an unusually large internal document part and was not processed."
                )
            names.add(info.filename)
    if "[Content_Types].xml" not in names or "word/document.xml" not in names:
        raise MalformedDocumentError("This file is a ZIP archive but not a Word (.docx) document.")



class DOCXParser:
    """
    DOCX text extraction parser using python-docx.
    Handles various Word document formats.
    """
    
    def __init__(self):
        """Initialize the DOCX parser."""
        pass
    
    def extract_text(self, docx_path: str) -> str:
        """
        Extract all text from a DOCX file.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Extracted text as string
            
        Raises:
            FileNotFoundError: If DOCX file doesn't exist
            Exception: For any extraction errors
        """
        text = ""
        
        try:
            # Check if file exists
            if not os.path.exists(docx_path):
                raise FileNotFoundError(f"DOCX file not found: {docx_path}")
            
            doc = Document(docx_path)
            logger.info(f"Processing DOCX file: {docx_path}")
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                para_text = para.text.strip()
                if para_text:
                    text += para_text + "\n"
            
            # Extract text from tables (common in resumes)
            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text:
                    text += "\n" + table_text + "\n"
            
            # Clean up the extracted text
            text = self._clean_text(text)
            
            logger.info(f"Total extracted text length: {len(text)} characters")
            return text
            
        except FileNotFoundError:
            logger.error(f"DOCX file not found: {docx_path}")
            raise
        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {str(e)}")
            raise
    
    def extract_text_strict(
        self,
        docx_path: str,
        *,
        max_uncompressed_bytes: int,
        max_chars: int = 300_000,
        max_entries: int = 1000,
        deadline: Optional[float] = None,
    ) -> str:
        """
        Extract text from an UNTRUSTED .docx with hard bounds (Task 5). Same text as extract_text()
        for a normal document, but: the archive is inspected first (decompression-bomb guard,
        encrypted / non-Word archives rejected), every failure is a typed DocumentError whose message
        is safe to show (no temp path, no library text), the amount of text is capped, and `deadline`
        (a time.monotonic() value) is checked between paragraphs and table rows.
        """
        inspect_docx_package(
            docx_path, max_uncompressed_bytes=max_uncompressed_bytes, max_entries=max_entries
        )

        def check_deadline():
            if deadline is not None and time.monotonic() > deadline:
                raise DocumentTimeoutError("Reading this DOCX took too long and was stopped.")

        def add(parts, chunk, total):
            total += len(chunk)
            if total > max_chars:
                raise DocumentTooComplexError(
                    "This document contains far more text than a resume should. "
                    "Please upload a shorter version."
                )
            parts.append(chunk)
            return total

        try:
            doc = Document(docx_path)
            parts, total = [], 0
            for para in doc.paragraphs:
                check_deadline()
                para_text = para.text.strip()
                if para_text:
                    total = add(parts, para_text + "\n", total)
            for table in doc.tables:
                check_deadline()
                rows = []
                for row in table.rows:
                    check_deadline()
                    row_text = [c for c in (cell.text.strip() for cell in row.cells) if c]
                    if row_text:
                        rows.append(" | ".join(row_text))
                if rows:
                    total = add(parts, "\n" + "\n".join(rows) + "\n" + "\n", total)
        except DocumentError:
            raise
        except MemoryError:
            raise
        except Exception as exc:
            raise MalformedDocumentError(_DOCX_DAMAGED, detail_for_log=f"{type(exc).__name__}: {exc}") from None

        text = self._clean_text("".join(parts))
        if not text:
            raise NoExtractableTextError("No readable text was found in this document.")
        logger.info(f"Extracted {len(text)} characters from DOCX (strict)")
        return text

    def _extract_table_text(self, table) -> str:
        """
        Extract text from a table, preserving structure.
        
        Args:
            table: python-docx table object
            
        Returns:
            Extracted table text
        """
        table_text = ""
        
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            
            if row_text:
                table_text += " | ".join(row_text) + "\n"
        
        return table_text
    
    def _clean_text(self, text: str) -> str:
        """
        Clean and normalize extracted text.
        
        Args:
            text: Raw extracted text
            
        Returns:
            Cleaned text
        """
        if not text:
            return ""
        
        lines = text.split('\n')
        cleaned_lines = []
        
        for line in lines:
            # Strip leading/trailing whitespace
            stripped = line.strip()
            
            # Skip empty lines
            if not stripped:
                continue
            
            # Remove multiple spaces
            cleaned = ' '.join(stripped.split())
            
            if cleaned:
                cleaned_lines.append(cleaned)
        
        # Join lines with proper spacing
        cleaned_text = '\n'.join(cleaned_lines)
        
        return cleaned_text
    
    def extract_metadata(self, docx_path: str) -> dict:
        """
        Extract DOCX metadata.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Dictionary containing metadata
        """
        metadata = {}
        
        try:
            doc = Document(docx_path)
            
            # Core properties
            core_props = doc.core_properties
            metadata = {
                'author': core_props.author,
                'title': core_props.title,
                'subject': core_props.subject,
                'keywords': core_props.keywords,
                'last_modified_by': core_props.last_modified_by,
                'revision': core_props.revision,
                'created': str(core_props.created),
                'modified': str(core_props.modified)
            }
            
        except Exception as e:
            logger.error(f"Error extracting metadata: {e}")
            metadata = {'error': str(e)}
        
        return metadata
    
    def extract_headers_footers(self, docx_path: str) -> dict:
        """
        Extract text from headers and footers.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Dictionary with header and footer text
        """
        headers_footers = {'headers': [], 'footers': []}
        
        try:
            doc = Document(docx_path)
            
            # Access section headers and footers
            for section in doc.sections:
                # Headers
                header = section.header
                for paragraph in header.paragraphs:
                    if paragraph.text.strip():
                        headers_footers['headers'].append(paragraph.text.strip())
                
                # Footers
                footer = section.footer
                for paragraph in footer.paragraphs:
                    if paragraph.text.strip():
                        headers_footers['footers'].append(paragraph.text.strip())
                        
        except Exception as e:
            logger.error(f"Error extracting headers/footers: {e}")
        
        return headers_footers


# Convenience function for quick extraction
def extract_text_from_docx(docx_path: str) -> str:
    """
    Quick text extraction from DOCX file.
    
    Args:
        docx_path: Path to the DOCX file
        
    Returns:
        Extracted text
    """
    parser = DOCXParser()
    return parser.extract_text(docx_path)
